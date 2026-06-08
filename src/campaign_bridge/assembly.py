"""Orchestrate §11 selections → validated Upload FB workbook."""
from __future__ import annotations

import os
import re
from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook

from config.paths import get_upload_fb_excel_path
from src.campaign_bridge.bundle import CampaignBundle
from src.campaign_bridge.open_rec_mapper import open_rec_to_plan_rows
from src.campaign_bridge.validation import resolve_account_sheet_name, validate_plan_rows
from src.upload_agent.build_upload_excel import (
    build_row,
    build_sheets,
    get_account_info,
    load_reference_data,
    sheet_name_safe,
)


def find_open_rec(
    report: dict[str, Any],
    country: str,
    account_label: str,
    strategy_type: int,
) -> dict[str, Any] | None:
    for section in report.get("countries") or []:
        if (section.get("country") or "").strip() != country.strip():
            continue
        recs = (section.get("open_recs") or {}).get(account_label) or []
        for rec in recs:
            if int(rec.get("strategy_type") or 0) == int(strategy_type):
                return rec
    return None


def build_plan_from_selection(
    report: dict[str, Any],
    *,
    country: str,
    account_label: str,
    selections: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[str]]:
    """selections: [{strategy_type, campaign_name, ad_titles: [str, str, str]}]."""
    plan_rows: list[dict[str, Any]] = []
    warnings: list[str] = []

    for sel in selections:
        st = int(sel["strategy_type"])
        rec = find_open_rec(report, country, account_label, st)
        if not rec:
            warnings.append(f"Strategy {st} no longer available — refresh report")
            continue
        plan_rows.extend(
            open_rec_to_plan_rows(
                rec,
                account_label=account_label,
                country=country,
                campaign_name=sel.get("campaign_name") or "",
                ad_titles=list(sel.get("ad_titles") or []),
            )
        )
    return plan_rows, warnings


def workbook_bytes_from_plan(
    ref: dict,
    account_label: str,
    plan_rows: list[dict[str, Any]],
) -> bytes:
    sheets = build_sheets(ref, {account_label: plan_rows})
    header_names = [h[1] for h in ref.get("template_headers") or []]
    wb = Workbook()
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    for account, rows in sheets.items():
        sheet_title = resolve_account_sheet_name(ref, account) or account
        ws = wb.create_sheet(title=sheet_name_safe(sheet_title))
        for col_idx, header in enumerate(header_names, start=1):
            ws.cell(row=1, column=col_idx, value=header)
        for row_idx, row_dict in enumerate(rows, start=2):
            for col_idx, header in enumerate(header_names, start=1):
                val = row_dict.get(header)
                ws.cell(row=row_idx, column=col_idx, value=val if val is not None else "")
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def assemble_account_upload(
    report: dict[str, Any],
    *,
    country: str,
    account_label: str,
    selections: list[dict[str, Any]],
    upload_fb_path: str | None = None,
) -> tuple[CampaignBundle | None, list[str], bytes | None, str | None]:
    """
    Validate and build workbook for one account Submit.

    Returns (bundle, errors, xlsx_bytes, download_filename).
    """
    plan_rows, lookup_warnings = build_plan_from_selection(
        report,
        country=country,
        account_label=account_label,
        selections=selections,
    )
    errors = list(lookup_warnings)

    excel_path = (upload_fb_path or "").strip() or get_upload_fb_excel_path()
    if not os.path.isfile(excel_path):
        errors.append(f"Upload FB workbook not found at {excel_path}")
        return None, errors, None, None

    ref = load_reference_data(excel_path)
    errors.extend(validate_plan_rows(ref, plan_rows, account_label=account_label))
    if errors:
        return None, errors, None, None

    sheet_name = resolve_account_sheet_name(ref, account_label) or account_label
    bundle = CampaignBundle(
        account_sheet=sheet_name,
        account_label=account_label,
        country=country,
        plan_rows=plan_rows,
    )
    xlsx_bytes = bundle.to_workbook_bytes(ref)
    slug = re.sub(r"[^a-zA-Z0-9]+", "_", account_label).strip("_") or "account"
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    filename = f"upload_{slug}_{stamp}.xlsx"
    return bundle, [], xlsx_bytes, filename
