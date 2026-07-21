"""Read the upload_plan.xlsx produced by Agent 8 (src/upload_agent) and track publish
status per row directly in the workbook — same "read/write by column name" approach as
build_upload_excel.py, so this stays in sync if the Agent 8 template gains/loses columns.

Idempotency (docs/agent10_fb_ads_playwright_upload.md §1.4): every row gets four status
columns (created on first load if missing). A row already marked "done" is skipped on
re-runs, so a crashed run can simply be restarted.
"""
from __future__ import annotations

import datetime
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

STATUS_COL = "Upload Status"
FB_CAMPAIGN_ID_COL = "FB Campaign ID"
UPLOADED_AT_COL = "Uploaded At"
ERROR_MESSAGE_COL = "Error Message"
STATUS_COLUMNS = (STATUS_COL, FB_CAMPAIGN_ID_COL, UPLOADED_AT_COL, ERROR_MESSAGE_COL)

STATUS_PENDING = "pending"
STATUS_DONE = "done"
STATUS_DRY_RUN_OK = "dry_run_ok"
STATUS_FAILED = "failed"


@dataclass
class PlanRow:
    sheet_name: str
    row_index: int  # 1-based Excel row number
    fields: dict[str, Any]
    status: str

    def get(self, *header_names: str, default: str = "") -> str:
        """Case-insensitive lookup across possible header spellings, e.g. get('Campaign Name', 'Campaign name')."""
        norm = {(h or "").strip().lower(): v for h, v in self.fields.items()}
        for name in header_names:
            v = norm.get((name or "").strip().lower())
            if v not in (None, ""):
                return str(v).strip()
        return default


def _normalize_header(h: Any) -> str:
    return str(h or "").strip()


def _header_index_map(ws) -> dict[str, int]:
    """header (trimmed, as-is case) -> 1-based column index, from row 1."""
    result = {}
    for col_idx, cell in enumerate(ws[1], start=1):
        h = _normalize_header(cell.value)
        if h:
            result[h] = col_idx
    return result


def _ensure_status_columns(ws, headers: dict[str, int]) -> dict[str, int]:
    """Append any missing status columns at the end of row 1. Returns updated header map."""
    next_col = max(headers.values(), default=0) + 1
    changed = False
    for col_name in STATUS_COLUMNS:
        if col_name not in headers:
            ws.cell(row=1, column=next_col, value=col_name)
            headers[col_name] = next_col
            next_col += 1
            changed = True
    return headers if not changed else headers


def load_pending_rows(plan_path: str | Path, account: str | None = None, limit: int | None = None) -> list[PlanRow]:
    """Read every non-empty row from every account sheet whose status isn't 'done' or 'dry_run_ok'
    (dry_run_ok rows are considered calibrated and are skipped too — pass status='pending' back
    manually, or delete the workbook's status cell, to re-include them)."""
    wb = load_workbook(plan_path)
    rows: list[PlanRow] = []
    for sheet_name in wb.sheetnames:
        if account and sheet_name.strip().lower() != account.strip().lower():
            continue
        ws = wb[sheet_name]
        headers = _header_index_map(ws)
        headers = _ensure_status_columns(ws, headers)
        status_col = headers[STATUS_COL]
        for row_idx in range(2, ws.max_row + 1):
            row_cells = ws[row_idx]
            if all(c.value in (None, "") for c in row_cells):
                continue
            fields = {}
            for h, col_idx in headers.items():
                if h in STATUS_COLUMNS:
                    continue
                fields[h] = ws.cell(row=row_idx, column=col_idx).value
            status = str(ws.cell(row=row_idx, column=status_col).value or "").strip().lower()
            if status in (STATUS_DONE, STATUS_DRY_RUN_OK):
                continue
            rows.append(PlanRow(sheet_name=sheet_name, row_index=row_idx, fields=fields, status=status or STATUS_PENDING))
            if limit and len(rows) >= limit:
                wb.save(plan_path)  # persist any newly-added status columns even on early exit
                return rows
    wb.save(plan_path)  # persist newly-added status columns
    return rows


def update_row_status(
    plan_path: str | Path,
    sheet_name: str,
    row_index: int,
    status: str,
    fb_campaign_id: str | None = None,
    error_message: str | None = None,
) -> None:
    """Write status back into the workbook immediately, so progress survives a crash mid-run."""
    wb = load_workbook(plan_path)
    ws = wb[sheet_name]
    headers = _header_index_map(ws)
    headers = _ensure_status_columns(ws, headers)
    ws.cell(row=row_index, column=headers[STATUS_COL], value=status)
    if fb_campaign_id is not None:
        ws.cell(row=row_index, column=headers[FB_CAMPAIGN_ID_COL], value=fb_campaign_id)
    if status in (STATUS_DONE, STATUS_DRY_RUN_OK):
        ws.cell(row=row_index, column=headers[UPLOADED_AT_COL], value=datetime.datetime.now().isoformat(timespec="seconds"))
    ws.cell(row=row_index, column=headers[ERROR_MESSAGE_COL], value=error_message or "")
    wb.save(plan_path)
