"""Архив статистики FB для дневника креативов.

Экспорт (`cursor example.xlsx`) — скользящее окно: он перезаписывается, и история
из него пропадает (в этой сессии файл на глазах ужался с 70 748 строк за два года
до 4 661 строки за четыре дня). Модуль накапливает все виденные строки в SQLite,
чтобы вердикты по связкам считались на полной истории.

Три таблицы:

* ``fact``         — текущее состояние строки. Ключ уникален в пределах экспорта
                     (проверено: 4661 ключ на 4661 строку), upsert по нему безопасен.
* ``revision``     — запись добавляется только когда метрики ключа изменились. По
                     ней измеряется лаг выручки: сколько дней проходит от открутки
                     до прихода депозита или возврата.
* ``snapshot_log`` — какой файл когда загружен и что принёс. Планировщик Windows
                     при сбое молчит, поэтому дыры должны быть видимы явно.

**Сентинел вместо NULL.** ``Ad Asset``/``Variant`` со значением ``-`` означают не
«нет оффера», а сбой джойна с трекером: 98.1% таких строк пусты по обоим полям
сразу, и 97.6% лежат на те же даты, что приписанные строки того же адсета. Хранить
их как SQL NULL нельзя — SQLite считает NULL-и различными и не обеспечивает
уникальность PRIMARY KEY, из-за чего ломается upsert. Поэтому в архиве это пустая
строка ``''`` (см. :data:`UNATTRIBUTED`), а разносит их по связкам уже
``creative_verdicts``.
"""
from __future__ import annotations

import datetime as dt
import os
import sqlite3
from dataclasses import dataclass, field
from typing import Iterable, Iterator, Sequence

import openpyxl

# ── схема экспорта ───────────────────────────────────────────────────────────

REQUIRED_COLUMNS = (
    "Country",
    "Account Alias",
    "Adset Name",
    "Ad Title",
    "# Leads",
    "# FTDs",
    "$ Spend",
    "CM",
    "Ad Asset",
    "$ Net",
    "Date",
)

#: Оффер/variant не сматчились с трекером. Не то же самое, что «оффера нет».
UNATTRIBUTED = ""

_NULLISH = {"", "-", "—", "none", "nan", "null"}

KEY_FIELDS = ("country", "account", "adset", "ad", "cm", "asset", "variant", "date")
METRIC_FIELDS = ("leads", "ftds", "spend", "net")

#: Порог, ниже которого расхождение денег считается округлением, а не правкой.
MONEY_EPS = 0.005


SCHEMA = """
CREATE TABLE IF NOT EXISTS fact (
    country   TEXT NOT NULL,
    account   TEXT NOT NULL,
    adset     TEXT NOT NULL,
    ad        TEXT NOT NULL,
    cm        TEXT NOT NULL,
    asset     TEXT NOT NULL,
    variant   TEXT NOT NULL,
    date      TEXT NOT NULL,
    leads     INTEGER NOT NULL DEFAULT 0,
    ftds      INTEGER NOT NULL DEFAULT 0,
    spend     REAL    NOT NULL DEFAULT 0,
    net       REAL    NOT NULL DEFAULT 0,
    first_seen TEXT NOT NULL,
    last_seen  TEXT NOT NULL,
    PRIMARY KEY (country, account, adset, ad, cm, asset, variant, date)
);

CREATE INDEX IF NOT EXISTS ix_fact_country_date ON fact (country, date);
CREATE INDEX IF NOT EXISTS ix_fact_adset        ON fact (country, adset);

CREATE TABLE IF NOT EXISTS revision (
    country   TEXT NOT NULL,
    account   TEXT NOT NULL,
    adset     TEXT NOT NULL,
    ad        TEXT NOT NULL,
    cm        TEXT NOT NULL,
    asset     TEXT NOT NULL,
    variant   TEXT NOT NULL,
    date      TEXT NOT NULL,
    observed_at TEXT NOT NULL,
    leads     INTEGER NOT NULL,
    ftds      INTEGER NOT NULL,
    spend     REAL    NOT NULL,
    net       REAL    NOT NULL
);

CREATE INDEX IF NOT EXISTS ix_revision_key
    ON revision (country, adset, ad, date, observed_at);

CREATE TABLE IF NOT EXISTS snapshot_log (
    id            INTEGER PRIMARY KEY AUTOINCREMENT,
    source_file   TEXT NOT NULL,
    source_mtime  TEXT NOT NULL,
    ingested_at   TEXT NOT NULL,
    observed_at   TEXT NOT NULL,
    rows_read     INTEGER NOT NULL,
    rows_inserted INTEGER NOT NULL,
    rows_updated  INTEGER NOT NULL,
    rows_unchanged INTEGER NOT NULL,
    date_min      TEXT,
    date_max      TEXT,
    UNIQUE (source_file, source_mtime)
);
"""

_KEY_JOIN = " AND ".join(f"f.{f} = i.{f}" for f in KEY_FIELDS)
_METRICS_DIFFER = (
    "f.leads <> i.leads OR f.ftds <> i.ftds"
    f" OR ABS(f.spend - i.spend) >= {MONEY_EPS} OR ABS(f.net - i.net) >= {MONEY_EPS}"
)


# ── нормализация ─────────────────────────────────────────────────────────────


def _clean_text(value) -> str:
    """Текстовое поле ключа. Пустое значение недопустимо — отдаём ``-``."""
    if value is None:
        return "-"
    return str(value).strip() or "-"


def _clean_attribution(value) -> str:
    """`Ad Asset` / `Variant`: ``-`` — сбой джойна, не значение."""
    if value is None:
        return UNATTRIBUTED
    s = str(value).strip()
    return UNATTRIBUTED if s.lower() in _NULLISH else s


def _clean_number(value) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace("$", "").replace(",", "")
    if not s or s.lower() in _NULLISH:
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0


def _clean_date(value) -> str | None:
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if value is None:
        return None
    s = str(value).strip()
    if not s or s.lower() in _NULLISH:
        return None
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%m/%d/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return dt.datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    return None


# ── чтение экспорта ──────────────────────────────────────────────────────────


class SchemaMismatch(RuntimeError):
    """Файл не является экспортом статистики (другой набор колонок)."""


@dataclass
class IngestResult:
    source_file: str
    rows_read: int = 0
    inserted: int = 0
    updated: int = 0
    unchanged: int = 0
    date_min: str | None = None
    date_max: str | None = None
    skipped: str | None = None

    @property
    def changed(self) -> int:
        return self.inserted + self.updated


def _column_map(header: Sequence) -> dict[str, int]:
    idx: dict[str, int] = {}
    for i, cell in enumerate(header):
        if cell is None:
            continue
        name = str(cell).strip()
        if name and name not in idx:
            idx[name] = i
    missing = [c for c in REQUIRED_COLUMNS if c not in idx]
    if missing:
        raise SchemaMismatch(f"нет колонок: {', '.join(missing)}")
    return idx


def read_export(path: str) -> Iterator[tuple]:
    """Пройти по строкам экспорта, отдавая нормализованные кортежи под ``fact``."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            raise SchemaMismatch("пустой лист") from None
        idx = _column_map(header)
        var_i = idx.get("Variant")

        def cell(row, name):
            i = idx[name]
            return row[i] if i < len(row) else None

        for row in rows:
            if not any(c is not None for c in row):
                continue
            date = _clean_date(cell(row, "Date"))
            if date is None:
                continue
            raw_variant = row[var_i] if var_i is not None and var_i < len(row) else None
            yield (
                _clean_text(cell(row, "Country")),
                _clean_text(cell(row, "Account Alias")),
                _clean_text(cell(row, "Adset Name")),
                _clean_text(cell(row, "Ad Title")),
                _clean_text(cell(row, "CM")),
                _clean_attribution(cell(row, "Ad Asset")),
                _clean_attribution(raw_variant),
                date,
                int(_clean_number(cell(row, "# Leads"))),
                int(_clean_number(cell(row, "# FTDs"))),
                _clean_number(cell(row, "$ Spend")),
                _clean_number(cell(row, "$ Net")),
            )
    finally:
        wb.close()


# ── запись в архив ───────────────────────────────────────────────────────────


def connect(db_path: str) -> sqlite3.Connection:
    parent = os.path.dirname(os.path.abspath(db_path))
    os.makedirs(parent, exist_ok=True)
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    conn.executescript(SCHEMA)
    return conn


def _already_ingested(conn: sqlite3.Connection, path: str, mtime: str) -> bool:
    cur = conn.execute(
        "SELECT 1 FROM snapshot_log WHERE source_file = ? AND source_mtime = ?",
        (os.path.basename(path), mtime),
    )
    return cur.fetchone() is not None


def ingest_file(
    conn: sqlite3.Connection,
    path: str,
    *,
    observed_at: str | None = None,
    skip_if_seen: bool = True,
) -> IngestResult:
    """Влить один экспорт в архив.

    ``observed_at`` — дата, на которую сделан снимок (по умолчанию mtime файла).
    Именно она позволяет измерить лаг выручки по таблице ``revision``, поэтому
    файлы обязательно заливать в хронологическом порядке (см. :func:`ingest_many`).

    Дедупликация и определение изменений идут внутри SQLite через временную
    таблицу — архив не помещается в память целиком.
    """
    mtime = dt.datetime.fromtimestamp(os.path.getmtime(path)).isoformat(timespec="seconds")
    observed = observed_at or mtime[:10]
    res = IngestResult(source_file=os.path.basename(path))

    if skip_if_seen and _already_ingested(conn, path, mtime):
        res.skipped = "уже загружен"
        return res

    try:
        records = list(read_export(path))
    except SchemaMismatch as exc:
        res.skipped = str(exc)
        return res
    except (OSError, KeyError, IndexError, ValueError) as exc:
        res.skipped = f"{type(exc).__name__}: {exc}"
        return res

    res.rows_read = len(records)
    if not records:
        res.skipped = "нет строк"
        return res

    dates = [r[7] for r in records]
    res.date_min, res.date_max = min(dates), max(dates)

    cols = ", ".join(KEY_FIELDS + METRIC_FIELDS)
    placeholders = ",".join("?" * (len(KEY_FIELDS) + len(METRIC_FIELDS)))

    conn.execute("DROP TABLE IF EXISTS temp.incoming")
    conn.execute(
        "CREATE TEMP TABLE incoming ("
        "country TEXT, account TEXT, adset TEXT, ad TEXT, cm TEXT,"
        "asset TEXT, variant TEXT, date TEXT,"
        "leads INTEGER, ftds INTEGER, spend REAL, net REAL,"
        "PRIMARY KEY (country, account, adset, ad, cm, asset, variant, date))"
    )
    # В пределах одного файла ключ уникален; ON CONFLICT — страховка от повтора.
    conn.executemany(
        f"INSERT INTO incoming ({cols}) VALUES ({placeholders})"
        " ON CONFLICT DO UPDATE SET leads=excluded.leads, ftds=excluded.ftds,"
        " spend=excluded.spend, net=excluded.net",
        records,
    )

    res.inserted = conn.execute(
        f"SELECT COUNT(*) n FROM incoming i LEFT JOIN fact f ON {_KEY_JOIN}"
        " WHERE f.date IS NULL"
    ).fetchone()["n"]
    res.updated = conn.execute(
        f"SELECT COUNT(*) n FROM incoming i JOIN fact f ON {_KEY_JOIN}"
        f" WHERE {_METRICS_DIFFER}"
    ).fetchone()["n"]
    res.unchanged = res.rows_read - res.inserted - res.updated

    conn.execute(
        f"INSERT INTO revision ({cols}, observed_at)"
        f" SELECT {', '.join('i.' + f for f in KEY_FIELDS + METRIC_FIELDS)}, ?"
        f" FROM incoming i LEFT JOIN fact f ON {_KEY_JOIN}"
        f" WHERE f.date IS NULL OR {_METRICS_DIFFER}",
        (observed,),
    )
    conn.execute(
        f"INSERT INTO fact ({cols}, first_seen, last_seen)"
        f" SELECT {cols}, ?, ? FROM incoming"
        " WHERE true"
        " ON CONFLICT (country, account, adset, ad, cm, asset, variant, date)"
        " DO UPDATE SET leads=excluded.leads, ftds=excluded.ftds,"
        " spend=excluded.spend, net=excluded.net, last_seen=excluded.last_seen"
        " WHERE fact.leads <> excluded.leads OR fact.ftds <> excluded.ftds"
        f" OR ABS(fact.spend - excluded.spend) >= {MONEY_EPS}"
        f" OR ABS(fact.net - excluded.net) >= {MONEY_EPS}",
        (observed, observed),
    )
    conn.execute("DROP TABLE temp.incoming")

    conn.execute(
        "INSERT OR REPLACE INTO snapshot_log (source_file, source_mtime, ingested_at,"
        " observed_at, rows_read, rows_inserted, rows_updated, rows_unchanged,"
        " date_min, date_max) VALUES (?,?,?,?,?,?,?,?,?,?)",
        (
            res.source_file,
            mtime,
            dt.datetime.now().isoformat(timespec="seconds"),
            observed,
            res.rows_read,
            res.inserted,
            res.updated,
            res.unchanged,
            res.date_min,
            res.date_max,
        ),
    )
    conn.commit()
    return res


def ingest_many(
    conn: sqlite3.Connection, paths: Iterable[str], *, verbose: bool = True
) -> list[IngestResult]:
    """Влить пачку экспортов в порядке mtime — от старых к новым.

    Порядок обязателен: ``revision`` должна отражать реальную хронологию
    наблюдений, иначе лаг выручки посчитается неверно.
    """
    ordered = sorted(paths, key=os.path.getmtime)
    results = []
    for path in ordered:
        res = ingest_file(conn, path)
        results.append(res)
        if verbose:
            if res.skipped:
                print(f"  ~ {res.source_file[:22]:22} пропуск: {res.skipped}")
            else:
                print(
                    f"  + {res.source_file[:22]:22} строк {res.rows_read:>6}"
                    f"  новых {res.inserted:>6}  обновлено {res.updated:>5}"
                    f"  {res.date_min}..{res.date_max}"
                )
    return results


# ── состояние архива ─────────────────────────────────────────────────────────


@dataclass
class ArchiveStatus:
    rows: int
    date_min: str | None
    date_max: str | None
    snapshots: int
    gaps: list[str] = field(default_factory=list)


def status(conn: sqlite3.Connection) -> ArchiveStatus:
    row = conn.execute("SELECT COUNT(*) n, MIN(date) lo, MAX(date) hi FROM fact").fetchone()
    snaps = conn.execute("SELECT COUNT(*) n FROM snapshot_log").fetchone()["n"]
    return ArchiveStatus(
        rows=row["n"],
        date_min=row["lo"],
        date_max=row["hi"],
        snapshots=snaps,
        gaps=missing_dates(conn),
    )


def missing_dates(conn: sqlite3.Connection, *, since: str | None = None) -> list[str]:
    """Календарные дни без единой строки — периоды, которым нельзя доверять."""
    row = conn.execute("SELECT MIN(date) lo, MAX(date) hi FROM fact").fetchone()
    if not row or not row["lo"]:
        return []
    lo = dt.date.fromisoformat(since or row["lo"])
    hi = dt.date.fromisoformat(row["hi"])
    have = {
        r["date"]
        for r in conn.execute(
            "SELECT DISTINCT date FROM fact WHERE date >= ?", (lo.isoformat(),)
        )
    }
    out, cur = [], lo
    while cur <= hi:
        if cur.isoformat() not in have:
            out.append(cur.isoformat())
        cur += dt.timedelta(days=1)
    return out
