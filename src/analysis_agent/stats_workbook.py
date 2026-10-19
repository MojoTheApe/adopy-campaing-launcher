"""Shared Stats workbook validation and loading (setup, dashboard settings, readers).

ACL-1864: заголовок — первая НЕпустая строка первого листа, а не обязательно
строка 1: форматирование может оставлять пустые строки над заголовками.
Пустыми считаются только пустые ячейки и строки из пробелов; ``0`` и ``False``
— содержимое. Если первая содержательная строка не валидная схема, дальше
НЕ ищем. Файл только читается, никогда не перезаписывается.
"""
from pathlib import Path
from typing import Any, Iterator, List, Optional, Tuple

STATS_COLUMNS = {
    "Country",
    "Account Alias",
    "Adset Name",
    "Ad Title",
    "# Leads",
    "# FTDs",
    "$ Spend",
    "CM",
    "Ad Asset",
    "$ Net",
}

# Порядок колонок, в котором их перечисляли читатели до общей схемы.
STATS_REQUIRED_COLUMNS: List[str] = [
    "Country", "Account Alias", "Adset Name", "Ad Title",
    "# Leads", "# FTDs", "$ Spend", "CM", "Ad Asset", "$ Net",
]


class StatsWorkbookError(ValueError):
    pass


def _is_blank(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return not value.strip()
    try:
        return value != value  # NaN
    except Exception:  # noqa: BLE001
        return False


def check_stats_headers(cells: List[Any]) -> List[str]:
    """Имена непустых заголовков первой содержательной строки; ValueError если схема не годится."""
    names = [str(v).strip() for v in cells if not _is_blank(v)]
    seen = set()
    dupes = set()
    for name in names:
        key = name.casefold()
        if key in seen:
            dupes.add(name)
        seen.add(key)
    if dupes:
        raise StatsWorkbookError("Stats workbook has duplicated columns: " + ", ".join(sorted(dupes)))
    missing = [c for c in STATS_REQUIRED_COLUMNS if c not in names]
    if missing:
        raise StatsWorkbookError(f"Missing columns in Excel: {missing}")
    return names


def first_content_row(rows: Iterator[Any]) -> Optional[Tuple[int, List[Any]]]:
    """(индекс, ячейки) первой непустой строки из итератора значений; None если все пусты."""
    for idx, row in enumerate(rows):
        row = list(row)
        if any(not _is_blank(v) for v in row):
            return idx, row
    return None


def find_stats_header_row(path: str) -> int:
    """Индекс (с 0) строки заголовков первого листа; валидирует схему."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        found = first_content_row(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if found is None:
        raise StatsWorkbookError("Stats workbook is empty: no header row found")
    check_stats_headers(found[1])
    return found[0]


def read_stats_workbook(path: str, usecols: Optional[List[str]] = None):
    """Прочитать Stats-книгу в DataFrame, заголовки — после ведущих пустых строк."""
    import pandas as pd

    header = find_stats_header_row(str(path))
    df = pd.read_excel(path, header=header)
    # Проверка схемы сравнивает имена без пробелов по краям — приводим колонки так же.
    df.columns = [c.strip() if isinstance(c, str) else c for c in df.columns]
    if usecols is not None:
        df = df[list(usecols)]
    return df


def validate_stats_workbook(path: str) -> str:
    workbook_path = Path(path)
    if not workbook_path.is_file():
        raise StatsWorkbookError("Stats workbook does not exist.")
    if workbook_path.suffix.casefold() not in {".xlsx", ".xlsm"}:
        raise StatsWorkbookError("Stats workbook must be an .xlsx or .xlsm file.")
    try:
        find_stats_header_row(str(workbook_path))
    except StatsWorkbookError as exc:
        raise StatsWorkbookError(f"Stats workbook is not usable: {exc}") from exc
    except Exception as exc:
        raise StatsWorkbookError(f"Stats workbook cannot be opened: {exc}") from exc
    return str(workbook_path.resolve())
