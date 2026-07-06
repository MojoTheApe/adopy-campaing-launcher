from __future__ import annotations

from dataclasses import dataclass, field
from typing import Any

from openpyxl import load_workbook


def _norm_header(val: Any) -> str:
    return str(val or "").strip().lower()


def _cell_str(val: Any) -> str:
    if val is None:
        return ""
    return str(val).strip()


@dataclass
class UploadRow:
    sheet_account: str
    row_number: int
    cells: dict[str, str] = field(default_factory=dict)

    def get(self, *header_names: str, default: str = "") -> str:
        for name in header_names:
            key = _norm_header(name)
            for h, v in self.cells.items():
                if _norm_header(h) == key:
                    return v
        return default


@dataclass
class UploadWorkbook:
    path: str
    sheets: dict[str, list[UploadRow]]


def read_upload_workbook(xlsx_path: str) -> UploadWorkbook:
    """Read Agent 8 output xlsx: one sheet per account, headers in row 1."""
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    sheets: dict[str, list[UploadRow]] = {}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        headers: list[str] = []
        rows: list[UploadRow] = []

        for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if r_idx == 1:
                headers = [_cell_str(c) for c in row if _cell_str(c)]
                continue
            if not headers:
                continue
            cells: dict[str, str] = {}
            has_data = False
            for col_idx, header in enumerate(headers):
                val = row[col_idx] if col_idx < len(row) else None
                s = _cell_str(val)
                if s:
                    has_data = True
                cells[header] = s
            if not has_data:
                continue
            if not _cell_str(cells.get("Campaign Name")) and not _cell_str(cells.get("Campaign name - country")):
                continue
            rows.append(UploadRow(sheet_account=sheet_name, row_number=r_idx, cells=cells))

        if rows:
            sheets[sheet_name] = rows

    wb.close()
    return UploadWorkbook(path=xlsx_path, sheets=sheets)
