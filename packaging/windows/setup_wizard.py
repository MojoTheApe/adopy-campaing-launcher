"""First-run Windows Setup Wizard, intentionally separate from ACL operations."""
from __future__ import annotations

import json
import os
import platform
import shutil
import socket
import tempfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Callable

from machine_config import (
    AccountAssignment,
    MachineConfiguration,
    MachineConfigError,
    apply_machine_configuration,
)


from src.analysis_agent.stats_workbook import (
    STATS_COLUMNS, StatsWorkbookError, validate_stats_workbook as _validate_stats_workbook,
)


class WizardValidationError(ValueError):
    pass


@dataclass(frozen=True)
class ReadinessCheck:
    label: str
    ok: bool
    detail: str


def setup_state_path(config_directory: str | Path) -> Path:
    return Path(config_directory) / "setup-state.json"


def setup_is_complete(config_directory: str | Path) -> bool:
    config_dir = Path(config_directory)
    try:
        state = json.loads(setup_state_path(config_dir).read_text(encoding="utf-8"))
        return bool(
            state.get("schema") == 1
            and state.get("status") == "ready"
            and (config_dir / ".env").is_file()
            and (config_dir / "ad_account_map.csv").is_file()
        )
    except Exception:
        return False


STATS_WORKBOOK_HELP = (
    "The Stats workbook is the Excel report your team lead keeps; it lists every campaign manager "
    "(CM column) with their ad accounts. ACL does not create or download it: ask your team lead or "
    "support to send you the current file, save it on this computer, then choose it here."
)


def default_stats_workbook() -> str:
    """The workbook in the usual location (FB_STATS_EXCEL or Desktop), or '' if it is not there."""
    try:
        from config.paths import get_excel_path

        candidate = get_excel_path()
    except Exception:
        return ""
    try:
        return candidate if candidate and Path(candidate).is_file() else ""
    except OSError:
        return ""


def validate_stats_workbook(path: str) -> str:
    if not (path or "").strip():
        raise WizardValidationError("No Stats workbook chosen. " + STATS_WORKBOOK_HELP)
    try:
        return _validate_stats_workbook(path)
    except StatsWorkbookError as exc:
        message = str(exc)
        if "does not exist" in message:
            message += " " + STATS_WORKBOOK_HELP
        else:
            message += " Ask your team lead for the current report file, then choose it again."
        raise WizardValidationError(message) from exc


def similar_names(name: str, candidates, limit: int = 3) -> list[str]:
    """Candidates that look like `name` (case-insensitive), best first."""
    import difflib

    by_fold: dict[str, str] = {}
    for candidate in candidates:
        by_fold.setdefault(candidate.casefold(), candidate)
    close = difflib.get_close_matches(name.casefold(), list(by_fold), n=limit, cutoff=0.6)
    return [by_fold[key] for key in close]


def default_creatives_directory() -> Path:
    """Documents\\ACL\\Creatives — created by setup, changeable later in the dashboard."""
    home = Path(os.environ.get("USERPROFILE") or Path.home())
    return home / "Documents" / "ACL" / "Creatives"


def validate_creatives_directory(path: str, *, create: bool = False) -> str:
    directory = Path(path) if (path or "").strip() else default_creatives_directory()
    if create and not directory.exists():
        try:
            directory.mkdir(parents=True, exist_ok=True)
        except OSError as exc:
            raise WizardValidationError(f"Creatives folder could not be created: {exc}") from exc
    if not directory.is_dir():
        raise WizardValidationError("Creatives folder does not exist.")
    if not os.access(directory, os.R_OK):
        raise WizardValidationError("Creatives folder is not readable.")
    return str(directory.resolve())


def validate_campaign_manager_name(value: str) -> str:
    name = (value or "").strip()
    if len(name) < 2:
        raise WizardValidationError("Enter your name exactly as it appears in the Stats workbook.")
    if len(name) > 120:
        raise WizardValidationError("Campaign manager name is too long.")
    if any(marker in name.casefold() for marker in ("campaign manager name", "your name", "example")):
        raise WizardValidationError("Replace the example campaign manager name with your own name.")
    return name


def build_manual_setup(*, campaign_manager_name: str, stats_workbook: str) -> MachineConfiguration:
    """Build machine configuration from fields entered in the setup wizard.

    Account labels come from the Stats workbook. Ad-account IDs and Multilogin
    profiles stay blank on purpose: on the Accounts screen the manager picks an
    open profile by name, and ACL reads the ad account from that profile.
    """
    from openpyxl import load_workbook

    name = validate_campaign_manager_name(campaign_manager_name)
    workbook = load_workbook(stats_workbook, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(rows)]
        by_name = {header.casefold(): index for index, header in enumerate(headers)}
        if "cm" not in by_name or "account alias" not in by_name:
            raise WizardValidationError("Stats workbook is missing CM or Account Alias columns.")
        cm_index = by_name["cm"]
        account_index = by_name["account alias"]
        aliases: dict[str, str] = {}
        known_managers: list[str] = []
        for row in rows:
            manager = str(row[cm_index] or "").strip() if cm_index < len(row) else ""
            if manager and manager not in known_managers:
                known_managers.append(manager)
            alias = str(row[account_index] or "").strip() if account_index < len(row) else ""
            if manager.casefold() == name.casefold() and alias:
                aliases.setdefault(alias.casefold(), alias)
    finally:
        workbook.close()

    if not aliases:
        close = similar_names(name, known_managers)
        hint = (
            "Did you mean: " + ", ".join(f"“{item}”" for item in close) + "? "
            "Go back to the name step and type it exactly as shown."
            if close
            else "No similar name was found. Check the spelling, or ask your team lead whether "
                 "you are listed in the CM column of this workbook."
        )
        raise WizardValidationError(
            f"“{name}” is not in the CM column of this Stats workbook, so the dashboard "
            f"would show no accounts. {hint}"
        )
    return MachineConfiguration(
        name,
        tuple(AccountAssignment(account_name=alias) for alias in sorted(aliases.values(), key=str.casefold)),
    )


def system_readiness(config_directory: str | Path, *, port: int = 9090) -> list[ReadinessCheck]:
    config_dir = Path(config_directory)
    checks: list[ReadinessCheck] = []
    windows_x64 = platform.system() == "Windows" and platform.machine().casefold() in {
        "amd64",
        "x86_64",
    }
    checks.append(ReadinessCheck("Windows x64", windows_x64, platform.platform()))

    try:
        config_dir.mkdir(parents=True, exist_ok=True)
        handle, temp_name = tempfile.mkstemp(prefix=".acl-write-test-", dir=config_dir)
        os.close(handle)
        Path(temp_name).unlink()
        writable = True
        writable_detail = str(config_dir)
    except Exception as exc:
        writable = False
        writable_detail = str(exc)
    checks.append(ReadinessCheck("Writable ACL data folder", writable, writable_detail))

    free_bytes = shutil.disk_usage(config_dir).free if config_dir.exists() else 0
    checks.append(
        ReadinessCheck(
            "Free disk space",
            free_bytes >= 500 * 1024 * 1024,
            f"{free_bytes // (1024 * 1024)} MB available",
        )
    )
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as sock:
        sock.settimeout(0.5)
        port_free = sock.connect_ex(("127.0.0.1", port)) != 0
    checks.append(
        ReadinessCheck(
            "Local dashboard port",
            port_free,
            f"port {port} is {'available' if port_free else 'already in use'}",
        )
    )
    return checks


#: Физические клавиши Windows: A, C, V, X. Смотреть надо именно на них.
#: Tk вешает вставку на латинскую «v», а при русской раскладке в событии
#: приходит «м» — правило не срабатывает, и Ctrl+V молча ничего не делает.
#: У команды раскладка русская, так что это не мелочь: поле почты нечем
#: заполнить, кроме как набрать руками (24.09).
CLIPBOARD_KEYCODES = {65: "<<SelectAll>>", 67: "<<Copy>>", 86: "<<Paste>>", 88: "<<Cut>>"}


def clipboard_event(keycode: int, state: int) -> str:
    """Какое действие с буфером просит это нажатие. Пусто — никакое.

    `state` — маска модификаторов Tk; бит 0x4 означает Control.
    """
    if not state & 0x4:
        return ""
    return CLIPBOARD_KEYCODES.get(keycode, "")


def enable_clipboard_shortcuts(widget, tk_module) -> None:
    """Вставка и копирование при любой раскладке, плюс меню правой кнопки."""

    def on_key(event):
        action = clipboard_event(getattr(event, "keycode", 0), getattr(event, "state", 0))
        if not action or not getattr(event, "widget", None):
            return None
        try:
            event.widget.event_generate(action)
        except Exception:
            return None
        return "break"

    widget.bind_all("<Control-KeyPress>", on_key, add="+")

    menu = tk_module.Menu(widget, tearoff=0)
    for label, action in (("Cut", "<<Cut>>"), ("Copy", "<<Copy>>"), ("Paste", "<<Paste>>")):
        menu.add_command(
            label=label,
            command=lambda a=action: menu.focus_get() and menu.focus_get().event_generate(a),
        )

    def on_right_click(event):
        try:
            event.widget.focus_set()
            menu.tk_popup(event.x_root, event.y_root)
        finally:
            menu.grab_release()

    widget.bind_all("<Button-3>", on_right_click, add="+")


def write_ready_state(config_directory: str | Path, version: str) -> None:
    path = setup_state_path(config_directory)
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "schema": 1,
        "status": "ready",
        "completed_at_utc": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
        "installer_version": version,
        "automatic_updates": False,
    }
    temp = path.with_name(f".{path.name}.tmp")
    temp.write_text(json.dumps(payload, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    os.replace(temp, path)


def sanitized_setup_summary(configuration: MachineConfiguration, version: str) -> str:
    """Return a support-safe completion summary with no identity or identifiers."""
    instances = sorted({assignment.multilogin_instance for assignment in configuration.accounts})
    return "\n".join(
        (
            "Adopy Campaign Manager setup summary",
            f"Installer version: {version}",
            "Status: base setup ready; account connection continues in the dashboard",
            f"Assigned accounts: {len(configuration.accounts)} (labels and IDs omitted)",
            f"Multilogin instances: {', '.join(instances)}",
            "Stats workbook and Creatives folder: validated (paths omitted)",
            "Safety mode: RPA / recorded_flow / attach-only",
            "Updates: offered in the dashboard, installed only when the user clicks Update",
            "Campaign actions during setup: none",
        )
    )


class SetupWizard:
    STEPS = (
        "Welcome",
        "System check",
        "Your name",
        "Local files",
        "Readiness",
        "Finish",
    )
    LAST = len(STEPS) - 1

    def __init__(self, root, config_directory: Path, version_text: str):
        import tkinter as tk

        self.tk = tk
        self.root = root
        self.config_directory = config_directory
        self.version_text = version_text
        self.step = 0
        self.configuration: MachineConfiguration | None = None
        self.completed = False
        self.campaign_manager_name = tk.StringVar(value=self._existing_env_value("MY_CM_NAME"))
        self.stats_path = tk.StringVar(
            value=self._existing_env_value("FB_STATS_EXCEL") or default_stats_workbook()
        )
        self.creatives_path = tk.StringVar(
            value=self._existing_env_value("FB_CREATIVES_ROOT") or str(default_creatives_directory()))
        self.status = tk.StringVar()

        root.title("Adopy Campaign Manager Setup")
        # 590 точек в высоту резали низ окна вместе с кнопкой «Next», и
        # человек оказывался в тупике на ровном месте (24.09). Запас оставлен
        # с тем же умыслом: шаг «Local files» ещё подрастёт (ACL-1305).
        root.geometry("820x680")
        root.minsize(760, 560)
        root.protocol("WM_DELETE_WINDOW", self._close)
        enable_clipboard_shortcuts(root, tk)
        self.header = tk.Frame(root, bg="#172036", padx=24, pady=18)
        self.header.pack(fill="x")
        self.heading = tk.Label(
            self.header,
            fg="white",
            bg="#172036",
            font=("Segoe UI", 18, "bold"),
        )
        self.heading.pack(anchor="w")
        self.progress = tk.Label(self.header, fg="#a9bbdf", bg="#172036", font=("Segoe UI", 10))
        self.progress.pack(anchor="w", pady=(4, 0))
        # Подвал кладётся ДО тела и прижимается к низу: при нехватке высоты
        # пакер вытесняет то, что уложено последним, — то есть ровно кнопки
        # «Next» и «Back». Теперь место под них отводится первым, и содержимое
        # шага ужимается само, но выйти со страницы можно всегда.
        self.footer = tk.Frame(root, padx=24, pady=16)
        self.footer.pack(side="bottom", fill="x")
        self.body = tk.Frame(root, padx=28, pady=24)
        self.body.pack(fill="both", expand=True)
        self.status_label = tk.Label(
            self.footer,
            textvariable=self.status,
            anchor="w",
            fg="#9b1c1c",
            wraplength=500,
            justify="left",
        )
        self.status_label.pack(side="left", fill="x", expand=True)
        self.back_button = tk.Button(self.footer, text="Back", width=10, command=self._back)
        self.back_button.pack(side="right", padx=(8, 0))
        self.next_button = tk.Button(self.footer, text="Next", width=12, command=self._next)
        self.next_button.pack(side="right")
        self._render()

    def _existing_env_value(self, key: str) -> str:
        path = self.config_directory / ".env"
        try:
            from dotenv import dotenv_values
            return str(dotenv_values(path, interpolate=False).get(key) or "")
        except Exception:
            return ""

    def _clear_body(self) -> None:
        for child in self.body.winfo_children():
            child.destroy()

    def _text(self, value: str, *, bold: bool = False, color: str = "#273248") -> None:
        self.tk.Label(
            self.body,
            text=value,
            anchor="w",
            justify="left",
            wraplength=700,
            fg=color,
            font=("Segoe UI", 11, "bold" if bold else "normal"),
        ).pack(fill="x", pady=(0, 12))

    def _render(self) -> None:
        self._clear_body()
        self.status.set("")
        self.heading.config(text=self.STEPS[self.step])
        self.progress.config(text=f"Step {self.step + 1} of {len(self.STEPS)}")
        self.back_button.config(state="normal" if 0 < self.step < self.LAST else "disabled")
        self.next_button.config(text="Close" if self.step == self.LAST else "Finish" if self.step == self.LAST - 1 else "Next")
        renderers: tuple[Callable[[], None], ...] = (
            self._welcome,
            self._system,
            self._identity,
            self._files,
            self._readiness,
            self._finish,
        )
        renderers[self.step]()
        self.next_button.config(state="normal")

    def _welcome(self) -> None:
        self._text("Set up Adopy Campaign Manager on this computer", bold=True)
        self._text(self.version_text)
        self._text(
            "This wizard only configures the local application. It cannot create, publish, or change a campaign."
        )
        self._text("New versions are offered in the dashboard and installed only when you click Update.")

    def _system(self) -> None:
        self.system_checks = system_readiness(self.config_directory)
        self._text("The installer checks this computer automatically; Python and Git are not prerequisites.")
        for check in self.system_checks:
            self._text(
                f"{'✓' if check.ok else '✕'}  {check.label} — {check.detail}",
                color="#17663a" if check.ok else "#9b1c1c",
            )

    def _identity(self) -> None:
        self._text("Enter your name exactly as it is written in the CM column of the Stats workbook.")
        self.tk.Label(self.body, text="Your name", anchor="w").pack(fill="x")
        self.tk.Entry(self.body, textvariable=self.campaign_manager_name).pack(fill="x", pady=(2, 12))
        self._text("ACL finds your ad accounts in the workbook by this name.")
        self._text("Run reports from this computer are labelled with this name, so support can tell whose "
                   "runs they are. You can change the label later in the dashboard (Accounts → This computer).")

    def _path_row(self, label: str, variable, directory: bool = False) -> None:
        from tkinter import filedialog

        self.tk.Label(self.body, text=label, anchor="w", font=("Segoe UI", 10, "bold")).pack(fill="x")
        row = self.tk.Frame(self.body)
        row.pack(fill="x", pady=(3, 12))
        self.tk.Entry(row, textvariable=variable).pack(side="left", fill="x", expand=True)

        def choose() -> None:
            selected = (
                filedialog.askdirectory()
                if directory
                else filedialog.askopenfilename(filetypes=[("Excel workbook", "*.xlsx *.xlsm")])
            )
            if selected:
                variable.set(selected)

        self.tk.Button(row, text="Browse…", command=choose).pack(side="left", padx=(8, 0))

    def _files(self) -> None:
        self._text("Use the file pickers; no path needs to be typed or edited.")
        self._text(STATS_WORKBOOK_HELP)
        if self.stats_path.get() and default_stats_workbook() == self.stats_path.get():
            self._text("A workbook was found in the usual place and is filled in below. "
                       "Change it only if you were sent a different file.", color="#17663a")
        self._path_row("Stats workbook (Excel)", self.stats_path)
        self._path_row("Creatives folder", self.creatives_path, directory=True)
        self._text("Put your videos and images into the Creatives folder. ACL creates it if it does not exist yet. "
                   "Both paths can be changed later in the dashboard.")

    def _readiness(self) -> None:
        checks = [
            ("Runtime", all(check.ok for check in getattr(self, "system_checks", []))),
            ("Your accounts", self.configuration is not None),
            ("Local files", bool(self.stats_path.get() and self.creatives_path.get())),
            ("Safety mode", True),
        ]
        self._text("The launcher will open only when every required group is ready.")
        for label, ok in checks:
            self._text(f"{'✓' if ok else '✕'}  {label}", color="#17663a" if ok else "#9b1c1c")
        self._text("Safety: RPA / recorded_flow / Multilogin attach-only", color="#17663a")
        self._text("•  Multilogin — connected in the dashboard, after setup", color="#8a6d1f")
        if self.configuration:
            self._text(f"Found {len(self.configuration.accounts)} account(s) for this campaign manager.")
            self._text("After the launcher opens, connect each account to its Multilogin profile on the Accounts screen.")
        self._text("Multilogin is connected in the dashboard, not here: open Accounts and press "
                   "“Connect Multilogin”, then sign in with your Multilogin email and password.")

    def _finish(self) -> None:
        self._text("Adopy Campaign Manager is ready on this computer.", bold=True, color="#17663a")
        self._text("No campaign action was performed. The first campaign run must remain supervised.")
        self._text("Click Close to open the dashboard. Then open Accounts and press Connect Multilogin "
                   "to sign in — until that is done the robot cannot open a profile.")
        self.tk.Button(
            self.body,
            text="Copy sanitized setup summary",
            command=self._copy_summary,
        ).pack(anchor="w", pady=10)

    def _copy_summary(self) -> None:
        if not self.configuration:
            return
        version = self.version_text.splitlines()[0].rsplit(" ", 1)[-1]
        self.root.clipboard_clear()
        self.root.clipboard_append(sanitized_setup_summary(self.configuration, version))
        self.status.set("Sanitized setup summary copied.")

    def _validate_step(self) -> None:
        if self.step == 1:
            failed = [check.label for check in self.system_checks if not check.ok]
            if failed:
                raise WizardValidationError("Resolve failed system checks: " + ", ".join(failed))
        elif self.step == 2:
            self.campaign_manager_name.set(validate_campaign_manager_name(self.campaign_manager_name.get()))
        elif self.step == 3:
            self.stats_path.set(validate_stats_workbook(self.stats_path.get()))
            self.creatives_path.set(validate_creatives_directory(self.creatives_path.get(), create=True))
            self.configuration = build_manual_setup(
                campaign_manager_name=self.campaign_manager_name.get(),
                stats_workbook=self.stats_path.get(),
            )
        elif self.step == 4:
            if not self.configuration:
                raise WizardValidationError("Go back and confirm your name and the Stats workbook.")
            apply_machine_configuration(
                self.configuration,
                self.config_directory,
                stats_workbook=self.stats_path.get(),
                creatives_directory=self.creatives_path.get(),
            )
            version = self.version_text.splitlines()[0].rsplit(" ", 1)[-1]
            write_ready_state(self.config_directory, version)
            self.completed = True

    def _next(self) -> None:
        if self.step == self.LAST:
            self._close()
            return
        try:
            self._validate_step()
        except (WizardValidationError, MachineConfigError, OSError) as exc:
            self.status.set(str(exc))
            return
        self.step += 1
        self._render()

    def _back(self) -> None:
        if self.step > 0:
            self.step -= 1
            self._render()

    def _close(self) -> None:
        self.root.destroy()


def run_setup_wizard(config_directory: str | Path, version_text: str) -> bool:
    try:
        import tkinter as tk

        root = tk.Tk()
    except Exception as exc:
        raise RuntimeError(f"Windows Setup Wizard could not start: {exc}") from exc
    wizard = SetupWizard(root, Path(config_directory), version_text)
    root.mainloop()
    return wizard.completed
