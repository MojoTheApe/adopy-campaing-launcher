"""
Marketing director Excel report (RU / EN):
- Read me sheet
- Summary by country
- Per-country: campaign managers → ad assets → top 2 combinations

Data source: FB stats Excel from config.paths.get_excel_path() (cursor example.xlsx by default).
Each build loads the latest calendar month in that file, takes every distinct Country value,
writes analysis_agent/_stats_doc_countries_lastmonth.xlsx (cache for other tools), and builds sheets
for all those geos — no hardcoded country list.

From repo root:
  python analysis_agent/build_marketing_director_report.py           # Russian only
  python analysis_agent/build_marketing_director_report.py --lang en # English only
  python analysis_agent/build_marketing_director_report.py --lang both
"""
from __future__ import annotations

import argparse
import json
import os
import sys
import time
from collections import Counter
from pathlib import Path
from typing import Literal

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# корень репозитория (родитель analysis_agent)
ROOT = Path(__file__).resolve().parents[1]
STATS_PATH = Path(__file__).resolve().parent / "_stats_doc_countries_lastmonth.xlsx"
OUT_PATH_RU = Path(__file__).resolve().parent / "Marketing_Director_Report.xlsx"
OUT_PATH_EN = Path(__file__).resolve().parent / "Marketing_Director_Report_EN.xlsx"

Lang = Literal["ru", "en"]

STRINGS: dict[str, dict[str, str]] = {
    "ru": {
        "sheet_readme": "Как читать",
        "sheet_summary": "Сводка",
        "doc_title": "Отчёт для директора по маркетингу",
        "period_unknown": "отчётный период",
        "readme_p1": "Период: {period}. Страны: все значения из колонки Country в исходном Excel за этот месяц ({n} гео). "
        "Источник — выгрузка статистики FB (документ из config / cursor example).",
        "readme_p2": "Лист «Сводка» — один ряд на страну: суммарные FTD, лиды, спенд, net, лучший ассет по FTD.",
        "readme_p3": "Листы стран: таблица 1 — кампейн-менеджеры; таблица 2 — ассеты; блок 3 — две лучшие связки "
        "(агрегат скрипта аналитики Strong/Interesting).",
        "readme_p4": "Важно: в исходных данных встречаются строки с Net при 0 FTD или 0 Spend — это корректировки/аллокации. "
        "Для оценки медиа опирайтесь на FTD, CR%, CPL, CPA.",
        "readme_data": "Файл данных: {name}",
        "sum_h_country": "Страна",
        "sum_h_ftd": "FTD всего",
        "sum_h_leads": "Лиды всего",
        "sum_h_spend": "Spend",
        "sum_h_net": "Net",
        "sum_h_cm": "Активных CM (с FTD>0)",
        "sum_h_top": "Топ-ассет по FTD",
        "sum_h_topftd": "FTD топ-ассета",
        "sec_cm": "Кампейн-менеджеры",
        "sec_asset": "По ассетам",
        "sec_combo": "Две лучшие связки (агрегат аналитики)",
        "combo_n": "Связка {n}",
        "no_rows": "Нет строк за период.",
        "combo_none": "Нет данных",
    },
    "en": {
        "sheet_readme": "Read me",
        "sheet_summary": "Summary",
        "doc_title": "Marketing director report",
        "period_unknown": "reporting period",
        "readme_p1": "Period: {period}. Countries: every distinct Country in the source Excel for that month ({n} geos). "
        "Source: Facebook stats file from config (cursor example).",
        "readme_p2": "Sheet «Summary»: one row per country — total FTDs, leads, spend, net, top ad asset by FTDs.",
        "readme_p3": "Country sheets: table 1 — campaign managers; table 2 — ad assets; section 3 — top two combinations "
        "(Strong/Interesting aggregate from the analytics script).",
        "readme_p4": "Note: some rows show Net with 0 FTD or $0 spend (adjustments / allocations). "
        "For media decisions, prioritise FTDs, CR%, CPL, and CPA.",
        "readme_data": "Data file: {name}",
        "sum_h_country": "Country",
        "sum_h_ftd": "Total FTDs",
        "sum_h_leads": "Total leads",
        "sum_h_spend": "Spend",
        "sum_h_net": "Net",
        "sum_h_cm": "CMs with FTD > 0",
        "sum_h_top": "Top ad asset (by FTDs)",
        "sum_h_topftd": "FTDs on top asset",
        "sec_cm": "Campaign managers",
        "sec_asset": "By ad asset",
        "sec_combo": "Top two combinations (analytics aggregate)",
        "combo_n": "Combination {n}",
        "no_rows": "No rows in this period.",
        "combo_none": "No data",
    },
}

# Top two combinations per country (same technical titles; descriptions localized)
TOP_COMBOS_RU: dict[str, list[tuple[str, str]]] = {
    "Bahrain": [
        (
            "GCC | all + dt_mixtest + 25+ + Vid dt1 - req | OTHER",
            "1 FTD, 8 лидов, CPL $22.45, CPA $179.62, CR 12.5% — лучший объём среди связок BH.",
        ),
        (
            "GCC | all + SM + 30 + 21222 | OTHER",
            "1 FTD, 1 лид, CPA $66.12 — сильный unit-CPA; клонировать после стабилизации Salik.",
        ),
    ],
    "Chile": [
        (
            "CL | all + IPO_VID + Video L | OTHER",
            "11 FTD, 266 лидов, CPL ~$29.8, CPA ~$720, CR ~4.1%.",
        ),
        (
            "CL | all + HV_VID + Video P | OTHER",
            "11 FTD, 293 лида — вторая по масштабу видеосвязка; масштаб вместе с Mercado Libre.",
        ),
    ],
    "Colombia": [
        (
            "CO | all + GMB + 1377 3 | OTHER",
            "17 FTD, 630 лидов — крупнейшая связка страны.",
        ),
        (
            "CO | all + GMB + 1377 2 | OTHER",
            "5 FTD, 279 лидов — вторая по FTD; проверить Bancolombia в Ads Manager по имени 1377.",
        ),
    ],
    "Ecuador": [
        (
            "EC | all + Trading_vid + 1319 3 | OTHER",
            "15 FTD, 977 лидов — якорь по Favorita/трейдингу.",
        ),
        (
            "EC | all + Trading_vid + 1389 2 | OTHER",
            "9 FTD, 408 лидов — вторая по объёму; держать в одной структуре с 1319.",
        ),
    ],
    "Guatemala": [
        (
            "GT | Trading_IOS + 442 Bancolombia - Copy | OTHER",
            "3 FTD, 180 лидов, низкий CPL/CPA в срезе — приоритет качества.",
        ),
        (
            "GT | PF_IOS + Por Vida Banco Industrial | OTHER",
            "1 FTD (малый n) — совпадает с оффером Banco Industrial; связать с рекомендациями PF_IOS в кабинете.",
        ),
    ],
    "Kuwait": [
        (
            "KW | EF + 30 + 1538_logov2 | OTHER",
            "1 FTD, 2 лида, CR 50% — узкий тест Finance Eagle.",
        ),
        (
            "KW | Reading_vid + — | OTHER",
            "1 FTD, 7 лидов — вторая структура; сверить спенд на уровне адсета в FB.",
        ),
    ],
    "Oman": [
        (
            "Нет строк Strong/Interesting в срезе отчёта",
            "Проверить выгрузку и название гео в столбце Country; при появлении данных — пересобрать аналитику.",
        ),
        (
            "Рекомендация",
            "Держать гео в отчёте для единого списка; заполнить связки после объёма в FB.",
        ),
    ],
    "Mexico": [
        (
            "MX | HV_img + CCBlack | OTHER",
            "15 FTD, 918 лидов — крупнейшая связка Amazon/MX в отчёте.",
        ),
        (
            "MX | PF/IOS/HV + Vid 700 - pro | PRO",
            "5 FTD, 224 лида, сильный CR ~2.2% — второй якорь масштаба.",
        ),
    ],
    "Qatar": [
        (
            "Нет строк Strong/Interesting в срезе отчёта",
            "Проверить выгрузку и название гео в столбце Country; при появлении данных — пересобрать аналитику.",
        ),
        (
            "Рекомендация",
            "Держать гео в отчёте для единого списка; заполнить связки после объёма в FB.",
        ),
    ],
    "United Arab Emirates": [
        (
            "UAE | Day + 1267 2 | OTHER",
            "2 FTD, 21 лид, CR ~9.5% — сильная конверсия лида в FTD.",
        ),
        (
            "UAE | Reading_vid + 1538 Logov2 | OTHER",
            "2 FTD, 72 лида — масштаб по лиду; связка Finance Eagle / reading.",
        ),
    ],
    "Uruguay": [
        (
            "Нет строк Strong/Interesting в срезе отчёта",
            "Проверить выгрузку и название гео в столбце Country; при появлении данных — пересобрать аналитику.",
        ),
        (
            "Рекомендация",
            "Держать гео в отчёте для единого списка; заполнить связки после объёма в FB.",
        ),
    ],
}

TOP_COMBOS_EN: dict[str, list[tuple[str, str]]] = {
    "Bahrain": [
        (
            "GCC | all + dt_mixtest + 25+ + Vid dt1 - req | OTHER",
            "1 FTD, 8 leads, CPL $22.45, CPA $179.62, CR 12.5% — best volume among BH combinations in the report.",
        ),
        (
            "GCC | all + SM + 30 + 21222 | OTHER",
            "1 FTD, 1 lead, CPA $66.12 — strong unit CPA on tiny n; clone after Salik structure is stable.",
        ),
    ],
    "Chile": [
        (
            "CL | all + IPO_VID + Video L | OTHER",
            "11 FTDs, 266 leads, CPL ~$29.8, CPA ~$720, CR ~4.1%.",
        ),
        (
            "CL | all + HV_VID + Video P | OTHER",
            "11 FTDs, 293 leads — second-largest video path; scale alongside Mercado Libre.",
        ),
    ],
    "Colombia": [
        (
            "CO | all + GMB + 1377 3 | OTHER",
            "17 FTDs, 630 leads — largest combination in the country.",
        ),
        (
            "CO | all + GMB + 1377 2 | OTHER",
            "5 FTDs, 279 leads — second by FTD; match Bancolombia in Ads Manager via 1377 naming.",
        ),
    ],
    "Ecuador": [
        (
            "EC | all + Trading_vid + 1319 3 | OTHER",
            "15 FTDs, 977 leads — core Favorita / trading angle.",
        ),
        (
            "EC | all + Trading_vid + 1389 2 | OTHER",
            "9 FTDs, 408 leads — second by volume; keep in one structure with 1319.",
        ),
    ],
    "Guatemala": [
        (
            "GT | Trading_IOS + 442 Bancolombia - Copy | OTHER",
            "3 FTDs, 180 leads, low CPL/CPA in this slice — quality-first priority.",
        ),
        (
            "GT | PF_IOS + Por Vida Banco Industrial | OTHER",
            "1 FTD (small n) — aligns with Banco Industrial; tie to PF_IOS recommendations in the ad account.",
        ),
    ],
    "Kuwait": [
        (
            "KW | EF + 30 + 1538_logov2 | OTHER",
            "1 FTD, 2 leads, CR 50% — narrow Finance Eagle test.",
        ),
        (
            "KW | Reading_vid + — | OTHER",
            "1 FTD, 7 leads — second structure; reconcile spend at ad-set level in Ads Manager.",
        ),
    ],
    "Mexico": [
        (
            "MX | HV_img + CCBlack | OTHER",
            "15 FTDs, 918 leads — largest Amazon/MX combination in the report.",
        ),
        (
            "MX | PF/IOS/HV + Vid 700 - pro | PRO",
            "5 FTDs, 224 leads, strong CR ~2.2% — second scale anchor.",
        ),
    ],
    "Oman": [
        (
            "No Strong/Interesting rows in this analytics slice",
            "Verify the export and the exact Country label; rebuild analytics after data appears.",
        ),
        (
            "Recommendation",
            "Keep the geo in the workbook for one list; fill combinations once FB volume is there.",
        ),
    ],
    "Qatar": [
        (
            "No Strong/Interesting rows in this analytics slice",
            "Verify the export and the exact Country label; rebuild analytics after data appears.",
        ),
        (
            "Recommendation",
            "Keep the geo in the workbook for one list; fill combinations once FB volume is there.",
        ),
    ],
    "United Arab Emirates": [
        (
            "UAE | Day + 1267 2 | OTHER",
            "2 FTDs, 21 leads, CR ~9.5% — strong lead-to-FTD conversion.",
        ),
        (
            "UAE | Reading_vid + 1538 Logov2 | OTHER",
            "2 FTDs, 72 leads — more lead volume; Finance Eagle / reading angle.",
        ),
    ],
    "Uruguay": [
        (
            "No Strong/Interesting rows in this analytics slice",
            "Verify the export and the exact Country label; rebuild analytics after data appears.",
        ),
        (
            "Recommendation",
            "Keep the geo in the workbook for one list; fill combinations once FB volume is there.",
        ),
    ],
}


def _combos_for(lang: Lang) -> dict[str, list[tuple[str, str]]]:
    return TOP_COMBOS_EN if lang == "en" else TOP_COMBOS_RU


HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F4E79")
SUB_FONT = Font(bold=True, size=12)
NOTE_FONT = Font(italic=True, size=10, color="404040")
THIN = Side(style="thin", color="B4B4B4")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")


def _last_month_slice(df: pd.DataFrame) -> pd.DataFrame:
    if "Date" not in df.columns:
        raise ValueError("Excel is missing Date column")
    if "Country" not in df.columns:
        raise ValueError("Excel is missing Country column")
    out = df.copy()
    out["Date"] = pd.to_datetime(out["Date"], errors="coerce")
    period_end = out["Date"].max()
    if pd.isna(period_end):
        raise ValueError("No valid dates in Excel")
    m = out["Date"].dt.to_period("M") == period_end.to_period("M")
    return out.loc[m]


def _countries_from_frame(df: pd.DataFrame) -> list[str]:
    s = df["Country"].astype(str).str.strip()
    s = s[(s.str.len() > 0) & ~s.str.lower().eq("nan")]
    return sorted(s.unique().tolist(), key=str.casefold)


def _country_sheet_title(country: str) -> str:
    """Excel worksheet name limit is 31 characters."""
    return country[:31]


def _load_source_stats() -> tuple[pd.DataFrame, list[str], Path]:
    """Read cursor-example stats Excel, keep latest month only, all countries in that slice."""
    sys.path.insert(0, str(ROOT))
    from config.paths import get_excel_path

    p = Path(get_excel_path())
    if not p.exists():
        raise FileNotFoundError(f"Stats Excel not found: {p}")
    raw = pd.read_excel(p)
    df = _last_month_slice(raw)
    if df.empty:
        raise ValueError("No rows in the latest month slice of the stats Excel.")
    countries = _countries_from_frame(df)
    if not countries:
        raise ValueError("No Country values in the latest month slice.")
    df.to_excel(STATS_PATH, index=False)
    return df, countries, p


def _style_header_row(ws, row: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def _write_table(
    ws,
    start_row: int,
    headers: list[str],
    data_rows: list[list],
    money_cols: set[int] | None = None,
) -> int:
    money_cols = money_cols or set()
    for j, h in enumerate(headers, start=1):
        ws.cell(row=start_row, column=j, value=h)
    _style_header_row(ws, start_row, len(headers))
    for i, row in enumerate(data_rows, start=1):
        r = start_row + i
        for j, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=j, value=val)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if j in money_cols and isinstance(val, (int, float)):
                cell.number_format = "#,##0.00"
    return start_row + len(data_rows) + 1


def _validate_country_sheets(wb: Workbook, countries: list[str]) -> None:
    """Every country sheet uses title = country[:31]; validate against that, not full Country string."""
    titles = [_country_sheet_title(c) for c in countries]
    dupes = [t for t, n in Counter(titles).items() if n > 1]
    if dupes:
        raise RuntimeError(
            "Country names truncate to duplicate Excel sheet names (31 char limit): "
            f"{sorted(set(dupes))}"
        )
    expected = set(titles)
    country_sheets = {n for n in wb.sheetnames if n in expected}
    if country_sheets != expected:
        missing = sorted(expected - country_sheets)
        extra = sorted(country_sheets - expected)
        raise RuntimeError(
            "Incomplete marketing director workbook: "
            f"missing sheets {missing}; extra {extra}; have {sorted(country_sheets)}"
        )


def _save_workbook_safe(wb: Workbook, output_path: Path, countries: list[str]) -> None:
    """Save to a temp file, then replace target (reduces risk of half-written xlsx)."""
    # #region agent log
    _dbg = Path(__file__).resolve().parent / "debug-5d3a16.log"
    # #endregion
    _validate_country_sheets(wb, countries)
    tmp = output_path.with_name(output_path.stem + "_tmp" + output_path.suffix)
    wb.save(tmp)
    try:
        os.replace(str(tmp), str(output_path))
    except OSError as e:
        # #region agent log
        try:
            payload = {
                "sessionId": "5d3a16",
                "timestamp": int(time.time() * 1000),
                "hypothesisId": "fix-atomic",
                "location": "build_marketing_director_report._save_workbook_safe",
                "message": "replace_failed_tmp_kept",
                "data": {"output": str(output_path), "tmp": str(tmp), "error": repr(e)},
                "runId": "post-fix",
            }
            with _dbg.open("a", encoding="utf-8") as f:
                f.write(json.dumps(payload, ensure_ascii=False) + "\n")
        except OSError:
            pass
        # #endregion
        raise RuntimeError(
            f"Could not replace {output_path} ({e!r}). "
            f"Close the workbook if it is open in Excel. "
            f"A complete file is available at: {tmp}"
        ) from e
    # #region agent log
    try:
        payload = {
            "sessionId": "5d3a16",
            "timestamp": int(time.time() * 1000),
            "hypothesisId": "fix-atomic",
            "location": "build_marketing_director_report._save_workbook_safe",
            "message": "save_ok",
            "data": {
                "output": str(output_path),
                "country_sheet_count": len(
                    [n for n in wb.sheetnames if n in {_country_sheet_title(c) for c in countries}]
                ),
            },
            "runId": "post-fix",
        }
        with _dbg.open("a", encoding="utf-8") as f:
            f.write(json.dumps(payload, ensure_ascii=False) + "\n")
    except OSError:
        pass
    # #endregion


def _autosize(ws, max_col: int, min_w: float = 10, max_w: float = 52) -> None:
    for col in range(1, max_col + 1):
        letter = get_column_letter(col)
        m = min_w
        for cell in ws[letter]:
            if cell.value is None:
                continue
            L = len(str(cell.value))
            m = min(max_w, max(m, L * 0.9 + 2))
        ws.column_dimensions[letter].width = m


def _write_combo_block(
    ws, row: int, S: dict[str, str], combo_pairs: list[tuple[str, str]]
) -> int:
    ws.cell(row=row, column=1, value=S["sec_combo"])
    ws.cell(row=row, column=1).font = SUB_FONT
    row += 1
    for idx, (title, desc) in enumerate(combo_pairs[:2], 1):
        ws.cell(row=row, column=1, value=S["combo_n"].format(n=idx))
        ws.cell(row=row, column=1).font = Font(bold=True)
        ws.cell(row=row, column=2, value=title)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=9)
        ws.cell(row=row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        row += 1
        ws.cell(row=row, column=2, value=desc)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=9)
        cell = ws.cell(row=row, column=2)
        cell.font = NOTE_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.fill = WARN_FILL
        row += 2
    return row


def build_workbook(lang: Lang = "ru", output_path: Path | None = None) -> Path:
    output_path = output_path or (OUT_PATH_EN if lang == "en" else OUT_PATH_RU)
    S = STRINGS[lang]
    TOP = _combos_for(lang)

    df, countries, src_path = _load_source_stats()
    for c in ["# Leads", "# FTDs", "$ Spend", "$ Net"]:
        df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0)
    df["CM"] = df["CM"].astype(str).str.strip()
    df["Ad Asset"] = df["Ad Asset"].fillna("—").astype(str).str.strip()

    df["Date"] = pd.to_datetime(df["Date"], errors="coerce")
    period_label = df["Date"].max()
    if pd.notna(period_label):
        period_str = period_label.strftime("%B %Y")
    else:
        period_str = S["period_unknown"]

    wb = Workbook()
    # --- Read me ---
    ws0 = wb.active
    ws0.title = S["sheet_readme"]
    ws0["A1"] = S["doc_title"]
    ws0["A1"].font = Font(bold=True, size=16, color="1F4E79")
    ws0["A3"] = S["readme_p1"].format(period=period_str, n=len(countries))
    ws0["A4"] = S["readme_p2"]
    ws0["A5"] = S["readme_p3"]
    ws0["A6"] = S["readme_p4"]
    ws0["A7"] = S["readme_data"].format(name=src_path.name)
    for r in range(3, 8):
        ws0.cell(row=r, column=1).alignment = Alignment(wrap_text=True)
    ws0.column_dimensions["A"].width = 100

    # --- Summary ---
    ws1 = wb.create_sheet(S["sheet_summary"], 1)
    sum_headers = [
        S["sum_h_country"],
        S["sum_h_ftd"],
        S["sum_h_leads"],
        S["sum_h_spend"],
        S["sum_h_net"],
        S["sum_h_cm"],
        S["sum_h_top"],
        S["sum_h_topftd"],
    ]
    summary_rows: list[list] = []
    for co in countries:
        sub = df[df["Country"].astype(str).str.strip() == co]
        if sub.empty:
            summary_rows.append([co, 0, 0, 0, 0, 0, "—", 0])
            continue
        ftd = int(sub["# FTDs"].sum())
        leads = int(sub["# Leads"].sum())
        spend = float(sub["$ Spend"].sum())
        net = float(sub["$ Net"].sum())
        cm_active = int((sub.groupby("CM")["# FTDs"].sum() > 0).sum())
        ast = (
            sub.groupby("Ad Asset", dropna=False)
            .agg(ftd=("# FTDs", "sum"))
            .reset_index()
            .sort_values("ftd", ascending=False)
        )
        top_a = str(ast.iloc[0]["Ad Asset"]) if len(ast) else "—"
        top_ftd = int(ast.iloc[0]["ftd"]) if len(ast) else 0
        summary_rows.append([co, ftd, leads, spend, net, cm_active, top_a, top_ftd])
    _write_table(ws1, 1, sum_headers, summary_rows, money_cols={4, 5})
    ws1.freeze_panes = "A2"
    if summary_rows:
        last_r = 1 + len(summary_rows)
        ws1.auto_filter.ref = f"A1:{get_column_letter(len(sum_headers))}{last_r}"
    _autosize(ws1, len(sum_headers))

    # --- Per country ---
    cm_cols = ["CM", "Deposits", "Leads", "Spend", "Net", "CPL", "CPA", "CR%", "ROI"]
    asset_cols = ["Ad Asset", "Deposits", "Leads", "Spend", "Net", "CPL", "CPA", "CR%", "ROI"]

    for co in countries:
        sheet_name = _country_sheet_title(co)
        ws = wb.create_sheet(sheet_name)
        row = 1
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        c = ws.cell(row=row, column=1, value=f"{co} — {period_str}")
        c.font = TITLE_FONT
        c.alignment = Alignment(horizontal="left", vertical="center")
        row += 2

        sub = df[df["Country"].astype(str).str.strip() == co]
        combos = TOP.get(co, [("—", S["combo_none"]), ("—", "")])
        if sub.empty:
            ws.cell(row=row, column=1, value=S["no_rows"])
            row += 1
            row += 2
            _write_combo_block(ws, row, S, combos)
            _autosize(ws, 9)
            continue

        ws.cell(row=row, column=1, value=S["sec_cm"])
        ws.cell(row=row, column=1).font = SUB_FONT
        row += 1
        cm_header_row = row

        cm = (
            sub.groupby("CM", dropna=False)
            .agg(Deposits=("# FTDs", "sum"), Leads=("# Leads", "sum"), Spend=("$ Spend", "sum"), Net=("$ Net", "sum"))
            .reset_index()
        )

        def fmt_cplpa(r):
            cpl = r["Spend"] / r["Leads"] if r["Leads"] > 0 else None
            cpa = r["Spend"] / r["Deposits"] if r["Deposits"] > 0 else None
            cr = round(100 * r["Deposits"] / r["Leads"], 2) if r["Leads"] > 0 else 0.0
            roi = (r["Net"] / r["Spend"]) if r["Spend"] > 0 else None
            return cpl, cpa, cr, roi

        cm_data = []
        for _, r in cm.sort_values("Deposits", ascending=False).iterrows():
            cpl, cpa, cr, roi = fmt_cplpa(r)
            cm_data.append(
                [
                    r["CM"],
                    int(r["Deposits"]),
                    int(r["Leads"]),
                    float(r["Spend"]),
                    float(r["Net"]),
                    round(cpl, 2) if cpl is not None else "—",
                    round(cpa, 2) if cpa is not None else "—",
                    cr,
                    round(roi, 2) if roi is not None else "—",
                ]
            )
        row = _write_table(ws, row, cm_cols, cm_data, money_cols={4, 5})
        ws.freeze_panes = ws.cell(row=cm_header_row + 1, column=1).coordinate
        row += 1

        ws.cell(row=row, column=1, value=S["sec_asset"])
        ws.cell(row=row, column=1).font = SUB_FONT
        row += 1

        ast = (
            sub.groupby("Ad Asset", dropna=False)
            .agg(Deposits=("# FTDs", "sum"), Leads=("# Leads", "sum"), Spend=("$ Spend", "sum"), Net=("$ Net", "sum"))
            .reset_index()
        )
        ast_data = []
        for _, r in ast.sort_values("Deposits", ascending=False).iterrows():
            cpl, cpa, cr, roi = fmt_cplpa(r)
            ast_data.append(
                [
                    r["Ad Asset"],
                    int(r["Deposits"]),
                    int(r["Leads"]),
                    float(r["Spend"]),
                    float(r["Net"]),
                    round(cpl, 2) if cpl is not None else "—",
                    round(cpa, 2) if cpa is not None else "—",
                    cr,
                    round(roi, 2) if roi is not None else "—",
                ]
            )
        row = _write_table(ws, row, asset_cols, ast_data, money_cols={4, 5})
        row += 2

        row = _write_combo_block(ws, row, S, combos)

        _autosize(ws, 9)

    _save_workbook_safe(wb, output_path, countries)
    return output_path


def main() -> None:
    ap = argparse.ArgumentParser(description="Build marketing director Excel report (RU/EN).")
    ap.add_argument(
        "--lang",
        choices=("ru", "en", "both"),
        default="ru",
        help="ru = Russian workbook only; en = English only; both = write both files.",
    )
    args = ap.parse_args()
    if args.lang in ("ru", "both"):
        print("Saved:", build_workbook("ru", OUT_PATH_RU))
    if args.lang in ("en", "both"):
        print("Saved:", build_workbook("en", OUT_PATH_EN))


if __name__ == "__main__":
    main()
