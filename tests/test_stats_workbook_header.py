"""ACL-1864: заголовки Stats-книги читаются после ведущих отформатированных пустых строк."""
from __future__ import annotations

import hashlib
from datetime import datetime
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from src.analysis_agent.morning_service import _load_enriched_df
from src.analysis_agent.stats_workbook import (
    STATS_REQUIRED_COLUMNS,
    StatsWorkbookError,
    read_stats_workbook,
    validate_stats_workbook,
)

HEADERS = STATS_REQUIRED_COLUMNS + ["Date", "Variant"]
ROWS = [
    ["Mexico", "MX-acc-1", "MX_trader_all_creative_trading_25", "Make money fast",
     100, 1, 50.0, "Evgeni Sedin", "trader_pro", 80.0, datetime(2026, 4, 18), "A"],
    ["Mexico", "MX-acc-1", "MX_trader_all_creative_trading_25", "Make money fast",
     0, 0, 0, "Evgeni Sedin", "trader_pro", 0, datetime(2026, 4, 17), "B"],
]


def _book(path: Path, *, lead_rows: int = 0, headers=None, rows=ROWS, prefix=None) -> Path:
    wb = Workbook()
    ws = wb.active
    fill = PatternFill("solid", fgColor="DDDDDD")
    for r in range(1, lead_rows + 1):  # отформатированные, но пустые
        for c in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill = fill
            cell.font = Font(bold=True)
            cell.number_format = "0.00"
    if prefix is not None:
        ws.cell(row=lead_rows + 1, column=1, value=prefix)
        lead_rows += 1
    if headers is not None:
        for c, h in enumerate(headers, 1):
            ws.cell(row=lead_rows + 1, column=c, value=h)
        for i, row in enumerate(rows):
            for c, v in enumerate(row, 1):
                ws.cell(row=lead_rows + 2 + i, column=c, value=v)
    wb.save(path)
    wb.close()
    return path


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_515_formatted_empty_rows_load_same_as_row1(tmp_path):
    plain = _book(tmp_path / "plain.xlsx", headers=HEADERS)
    shifted = _book(tmp_path / "shifted.xlsx", lead_rows=515, headers=HEADERS)
    before = _sha(shifted)
    a = _load_enriched_df(str(plain))
    b = _load_enriched_df(str(shifted))
    pd.testing.assert_frame_equal(a, b)
    assert len(b) == 2
    assert validate_stats_workbook(str(shifted))
    assert _sha(shifted) == before  # книга не переписана


def test_row1_workbook_with_extra_columns_and_reordering_still_loads(tmp_path):
    order = list(reversed(HEADERS))
    rows = [list(reversed(r)) for r in ROWS]
    book = _book(tmp_path / "rev.xlsx", headers=order, rows=rows)
    df = read_stats_workbook(str(book))
    assert list(df.columns) == order
    assert df["# Leads"].tolist() == [100, 0]


@pytest.mark.parametrize("bad", [HEADERS[:5], [h for h in HEADERS if h != "CM"]])
def test_partial_headers_rejected(tmp_path, bad):
    book = _book(tmp_path / "bad.xlsx", lead_rows=3, headers=bad, rows=[])
    with pytest.raises(StatsWorkbookError):
        read_stats_workbook(str(book))
    with pytest.raises(StatsWorkbookError):
        validate_stats_workbook(str(book))


def test_duplicate_normalized_header_rejected(tmp_path):
    book = _book(tmp_path / "dup.xlsx", lead_rows=2, headers=HEADERS + [" country "], rows=[])
    with pytest.raises(StatsWorkbookError, match="duplicated"):
        read_stats_workbook(str(book))


@pytest.mark.parametrize("prefix", [0, False])
def test_zero_and_false_are_content_so_no_later_header_search(tmp_path, prefix):
    book = _book(tmp_path / "z.xlsx", lead_rows=4, headers=HEADERS, prefix=prefix)
    with pytest.raises(StatsWorkbookError):
        read_stats_workbook(str(book))
    with pytest.raises(StatsWorkbookError):
        validate_stats_workbook(str(book))


def test_whitespace_only_cells_are_blank(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "   "
    for c, h in enumerate(HEADERS, 1):
        ws.cell(row=2, column=c, value=h)
    path = tmp_path / "ws.xlsx"
    wb.save(path)
    wb.close()
    assert len(read_stats_workbook(str(path))) == 0
    assert validate_stats_workbook(str(path))


def test_empty_sheet_rejected(tmp_path):
    book = _book(tmp_path / "empty.xlsx", lead_rows=10)
    with pytest.raises(StatsWorkbookError, match="empty"):
        read_stats_workbook(str(book))
    with pytest.raises(StatsWorkbookError):
        validate_stats_workbook(str(book))


def test_morning_report_builder_accepts_shifted_headers(tmp_path):
    from src.analysis_agent.morning_report import build_morning_report

    plain = _book(tmp_path / "plain.xlsx", headers=HEADERS)
    shifted = _book(tmp_path / "shifted.xlsx", lead_rows=515, headers=HEADERS)
    ra = build_morning_report(excel_path=str(plain), my_cm="Evgeni Sedin")
    rb = build_morning_report(excel_path=str(shifted), my_cm="Evgeni Sedin")
    ra.pop("source_excel", None)
    rb.pop("source_excel", None)
    assert ra.get("countries") == rb.get("countries")
