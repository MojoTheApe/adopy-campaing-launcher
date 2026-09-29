"""ACL-1342 (аудит безопасности, S10): офер с именем `=1+1` не должен стать
формулой в скачанном Excel.

`/export/assets` пишет `Country`/`Asset` прямо из отчёта — данных, которые
владелец не контролирует (их вводят байеры при заведении офера). openpyxl
сам решает, что ячейка — формула, по одному признаку: строковое значение
начинается с `=`. Открыв такой экспорт, Excel посчитает `=1+1` как `2`,
а какой-нибудь `=HYPERLINK(...)` или `=cmd|...` — выполнит хуже.

Проверяется настоящий эндпоинт (`web.export_assets`), а не пересборка
Excel-файла вручную: фикс мог остаться только в тестовой копии кода.
"""
from __future__ import annotations

import asyncio
import io
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from src.web import app as web


def _report_with_offer(country: str, offer_name: str) -> dict:
    return {
        "countries": [
            {
                "country": country,
                "asset_summary": [
                    {"offer": offer_name, "leads": 10, "deposits": 2, "net": 30.0},
                ],
            }
        ]
    }


def _load_export(monkeypatch, report: dict):
    monkeypatch.setattr(web, "get_cached_report", lambda: (report, None, []))
    response = asyncio.run(web.export_assets())
    wb = openpyxl.load_workbook(io.BytesIO(response.body))
    return wb.active


@pytest.mark.parametrize("offer_name", ["=1+1", "+1+1", "-2+2", "@SUM(1,1)"])
def test_a_formula_looking_offer_name_round_trips_as_literal_text(monkeypatch, offer_name):
    ws = _load_export(monkeypatch, _report_with_offer("Mexico", offer_name))

    cell = ws.cell(row=2, column=2)  # столбец "Asset"
    assert cell.value == offer_name, "значение исказилось при экспорте"
    assert cell.data_type == "s", (
        "ячейка не помечена текстовой — Excel посчитает офер формулой"
    )


def test_an_ordinary_offer_name_is_unaffected(monkeypatch):
    ws = _load_export(monkeypatch, _report_with_offer("Mexico", "Casino Welcome Bonus"))

    country_cell = ws.cell(row=2, column=1)
    offer_cell = ws.cell(row=2, column=2)
    assert country_cell.value == "Mexico"
    assert offer_cell.value == "Casino Welcome Bonus"
    assert offer_cell.data_type == "s"


def test_numeric_columns_stay_numbers(monkeypatch):
    """Фикс трогает только текстовые столбцы 1–2, а не Leads/Deposits/CR/Value."""
    ws = _load_export(monkeypatch, _report_with_offer("Mexico", "=1+1"))

    leads_cell = ws.cell(row=2, column=3)
    assert leads_cell.value == 10
    assert leads_cell.data_type == "n"
