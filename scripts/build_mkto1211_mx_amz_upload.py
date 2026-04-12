#!/usr/bin/env python3
"""Build FB bulk-upload Excel for two MX Amazon campaigns on FacebookMKTO1211 Latam Amz."""
from __future__ import annotations

import sys
import shutil
from pathlib import Path

from openpyxl import Workbook, load_workbook

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parents[1]
REF = Path(r"C:\Users\EvgeniSedin\Desktop\Upluad FB.xlsx")
OUT_SLIM = ROOT / "analysis_agent" / "FacebookMKTO1211_LatamAmz_MX_Amazon_2campaigns_upload.xlsx"
OUT_FULL = ROOT / "analysis_agent" / "FacebookMKTO1211_LatamAmz_MX_Amazon_2campaigns_FULL.xlsx"
UPLOAD_SHEET = "FBMKTO1211_MX_2campaigns"

ACCOUNT = "FacebookMKTO1211 Latam Amz"
PIXEL = "1448065443673419"
LINK = (
    "https://trk.coralora.com/Swj5gkf3?fbclid={fbclid}&utm_campaign={{campaign.name}}"
    "&utm_source={{site_source_name}}&utm_placement={{placement}}&campaign_id={{campaign.id}}"
    "&adset_id={{adset.id}}&ad_id={{ad.id}}&adset_name={{adset.name}}&ad_name={{ad.name}}"
)

CAMPAIGN_VIDEO = "MX_AMZ_All_Video700_SM_IOS_HV_40_25.02"
CAMPAIGN_IMAGE = "MX_AMZN_all_Img700_PF_40_20.03"


def _headers(ws) -> list:
    return [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]


def _col_idx(headers: list, name: str) -> int | None:
    for i, h in enumerate(headers):
        if h == name:
            return i + 1
    return None


def _get_json_for_targeting(wb, name: str) -> str:
    ws = wb["Interests"]
    for r in range(2, ws.max_row + 1):
        n = ws.cell(r, 1).value
        if n is not None and str(n).strip() == name:
            j = ws.cell(r, 2).value
            return str(j) if j is not None else ""
    return ""


def main() -> None:
    if not REF.is_file():
        raise SystemExit(f"Reference Excel not found: {REF}")

    wb = load_workbook(REF, data_only=True)
    src = wb["MX Amz"]
    headers = _headers(src)
    max_c = src.max_column
    base = [src.cell(2, c).value for c in range(1, max_c + 1)]

    hv_json = _get_json_for_targeting(wb, "HV_MX")
    pf_json = _get_json_for_targeting(wb, "PF")
    wb.close()

    def setv(row: list, header: str, value) -> None:
        idx = _col_idx(headers, header)
        if idx is not None:
            row[idx - 1] = value

    # Видео: три разных файла из справочника (11163/21163/31163); в шаблоне MX Amz часто один mp4 на три копирайта —
    # здесь отдельные ролики под reg / pro / simple2.
    video_rows = [
        ("1163", "11163.mp4", "11163_Vid 2 - reg", "ES Invest Like a pro Amazon"),
        ("1163", "21163.mp4", "21163_Vid 2 - pro", "ES Small investment  Amazon"),
        ("1163", "31163.mp4", "31163_Vid 2 - simple2", "ES Amazon Every package"),
    ]

    image_rows = [
        ("134", "134.jpg", "134_700 1", "ES Invest Like a pro Amazon"),
        ("1134", "1134.jpg", "1134_700 2", "ES Small investment  Amazon"),
        ("10422", "10422.jpg", "10422_Blurry 2", "ES Amazon Every package"),
    ]

    out_rows: list[list] = []

    for cr_num, vid_fn, ad_name, copy_name in video_rows:
        row = list(base)
        setv(row, "Pixel", PIXEL)
        setv(row, "Link", LINK)
        setv(row, "Optimized Conversion Tracking Pixels", f"tp:{PIXEL}")
        setv(row, "Pixel id", PIXEL)
        setv(row, "Campaign name - country", "MX")
        setv(row, "Asset", "AMZ")
        setv(row, "OS", "All")
        setv(row, "Creative", "Video")
        setv(row, "Targeting", "HV_MX")
        setv(row, "Age", 40)
        setv(row, "Age Min", 40)
        setv(row, "Creative number", int(cr_num) if str(cr_num).isdigit() else cr_num)
        setv(row, "Date", "25.02")
        setv(row, "Copy", copy_name)
        setv(row, "Campaign Name", CAMPAIGN_VIDEO)
        setv(row, "Ad Set Name", CAMPAIGN_VIDEO)
        setv(row, "Flexible Inclusions", hv_json)
        setv(row, "Ad Name", ad_name)
        setv(row, "Ad Draft", vid_fn)
        setv(row, "Video File Name", vid_fn)
        setv(row, "Creative Type", "Video")
        setv(row, "Image File name", None)
        setv(row, "Video ID", None)
        setv(row, "Image Hash", None)
        out_rows.append(row)

    for img_num, img_fn, ad_name, copy_name in image_rows:
        row = list(base)
        setv(row, "Pixel", PIXEL)
        setv(row, "Link", LINK)
        setv(row, "Optimized Conversion Tracking Pixels", f"tp:{PIXEL}")
        setv(row, "Pixel id", PIXEL)
        setv(row, "Campaign name - country", "MX")
        setv(row, "Asset", "AMZ")
        setv(row, "OS", "All")
        setv(row, "Creative", "Image")
        setv(row, "Targeting", "PF")
        setv(row, "Age", 40)
        setv(row, "Age Min", 40)
        setv(row, "Creative number", int(img_num) if str(img_num).isdigit() else img_num)
        setv(row, "Date", "20.03")
        setv(row, "Copy", copy_name)
        setv(row, "Campaign Name", CAMPAIGN_IMAGE)
        setv(row, "Ad Set Name", CAMPAIGN_IMAGE)
        setv(row, "Flexible Inclusions", pf_json)
        setv(row, "Ad Name", ad_name)
        setv(row, "Ad Draft", img_fn)
        setv(row, "Image File name", img_fn)
        setv(row, "Video File Name", None)
        setv(row, "Video ID", None)
        setv(row, "Creative Type", "Image")
        out_rows.append(row)

    def write_upload_sheet(wb) -> None:
        if UPLOAD_SHEET in wb.sheetnames:
            wb.remove(wb[UPLOAD_SHEET])
        ws_out = wb.create_sheet(title=UPLOAD_SHEET)
        for c in range(1, max_c + 1):
            ws_out.cell(1, c, value=headers[c - 1])
        for r_i, row in enumerate(out_rows, start=2):
            for c_i, val in enumerate(row, start=1):
                ws_out.cell(r_i, c_i, value=val)

    OUT_SLIM.parent.mkdir(parents=True, exist_ok=True)

    slim = Workbook()
    slim.remove(slim.active)
    ws_slim = slim.create_sheet(title=UPLOAD_SHEET)
    for c in range(1, max_c + 1):
        ws_slim.cell(1, c, value=headers[c - 1])
    for r_i, row in enumerate(out_rows, start=2):
        for c_i, val in enumerate(row, start=1):
            ws_slim.cell(r_i, c_i, value=val)
    meta = slim.create_sheet(title="_Readme")
    meta["A1"] = "Ads Manager account"
    meta["B1"] = ACCOUNT
    meta["A2"] = "Pixel"
    meta["B2"] = PIXEL
    meta["A3"] = "Link (tracking)"
    meta["B3"] = LINK
    meta["A5"] = (
        "Лист склонирован с MX Amz (Upluad FB.xlsx): бюджет, плейсменты, excluded countries, "
        "attribution, optimization и т.д. Pixel/Link — из листа Accounts для этого аккаунта."
    )
    meta["A6"] = (
        "Видео-кампания: Flexible Inclusions = HV_MX (в справочнике Interests нет отдельного SM/IOS/HV). "
        "Файлы 11163/21163/31163.mp4 — из листа Video file name; при необходимости замените на реальные «Vid 2»."
    )
    meta["A7"] = (
        "Картинки: id 134, 1134, 10422 — из Image FIle name; подставьте свои файлы для 700/Blurry, если другие id в библиотеке."
    )
    meta["A9"] = "Полная копия со справочниками:"
    meta["B9"] = str(OUT_FULL.name)

    slim.save(OUT_SLIM)
    print(f"Written (upload-only): {OUT_SLIM}")

    shutil.copy2(REF, OUT_FULL)
    full_wb = load_workbook(OUT_FULL)
    write_upload_sheet(full_wb)
    rm = full_wb.create_sheet(title="_Readme_MKTO1211")
    rm["A1"] = "Импорт в FB: лист"
    rm["B1"] = UPLOAD_SHEET
    rm["A2"] = "Аккаунт"
    rm["B2"] = ACCOUNT
    rm["A4"] = (
        "Остальные листы — полная копия Upluad FB.xlsx (Accounts, Ad copy, Countries, Interests, "
        "Video/Image file name и т.д.)."
    )
    full_wb.save(OUT_FULL)
    full_wb.close()
    print(f"Written (full + refs): {OUT_FULL}")


if __name__ == "__main__":
    main()
