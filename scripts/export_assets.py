"""Generate output/assets_YYYY-MM-DD.xlsx with asset summary per country."""
from __future__ import annotations

import io
import sys
from datetime import datetime
from pathlib import Path

_ROOT = Path(__file__).resolve().parents[1]
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from src.web.cache import get_cached_report


def build_assets_excel(report: dict) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Assets"

    hdr_fill = PatternFill("solid", fgColor="1F2937")
    hdr_font = Font(bold=True, color="58A6FF", size=10)
    hdr_align = Alignment(horizontal="center", vertical="center")

    headers = ["Country", "Asset", "Leads", "Deposits", "CR %", "Player Value ($)"]
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = hdr_align

    alt_fill = PatternFill("solid", fgColor="161B22")
    num_align = Alignment(horizontal="right")

    row_idx = 2
    for c in report.get("countries", []):
        country = c.get("country", "")
        for asset in c.get("asset_summary", []):
            leads = asset.get("leads", 0)
            deposits = asset.get("deposits", 0)
            cr = round(deposits / leads * 100, 1) if leads > 0 else 0.0
            net = asset.get("net", 0.0)

            fill = alt_fill if row_idx % 2 == 0 else None
            row_data = [country, asset.get("offer", ""), leads, deposits, cr, net]
            for col_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                if fill:
                    cell.fill = fill
                if col_idx >= 3:
                    cell.alignment = num_align
                if col_idx == 5:
                    cell.number_format = '0.0"%"'
                if col_idx == 6:
                    cell.number_format = '#,##0.00'
            row_idx += 1

    col_widths = [16, 26, 10, 12, 10, 18]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{row_idx - 1}"
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def main() -> None:
    print("Loading report from cache...")
    report, _, _ = get_cached_report()

    countries = report.get("countries", [])
    total_rows = sum(len(c.get("asset_summary", [])) for c in countries)
    print(f"Found {len(countries)} countries, {total_rows} asset rows")

    xlsx_bytes = build_assets_excel(report)

    out_dir = _ROOT / "output"
    out_dir.mkdir(exist_ok=True)
    date_str = datetime.now().strftime("%Y-%m-%d")
    out_path = out_dir / f"assets_{date_str}.xlsx"
    out_path.write_bytes(xlsx_bytes)

    print(f"Saved: {out_path}")

    for c in countries:
        assets = c.get("asset_summary", [])
        if assets:
            print(f"  {c['country']}: {len(assets)} assets")


if __name__ == "__main__":
    main()
